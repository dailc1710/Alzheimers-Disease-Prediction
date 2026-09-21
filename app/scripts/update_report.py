"""Refresh the supplied V3 report from the current active metadata and audit tables."""

from __future__ import annotations

import argparse
import json
import os
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[1]
WORKSPACE_ROOT = ROOT.parent
DELIVERABLES_ROOT = WORKSPACE_ROOT / "deliverables"
DEFAULT_REPORT = DELIVERABLES_ROOT / "reports" / "I1_2510_E0_-_Nhom_1-_Alzheimer_project_report_V3_revised.docx"
DEFAULT_OUTPUT = DELIVERABLES_ROOT / "reports" / "I1_2510_E0_-_Nhom_1-_Alzheimer_project_report_V3_aligned.docx"


def _decimal(value: object, digits: int = 4) -> str:
    """Format a numeric report value with the document's decimal-comma style."""

    if value is None or pd.isna(value):
        return "—"
    return f"{float(value):.{digits}f}".replace(".", ",")


def _integer(value: object) -> str:
    return f"{int(value):,}".replace(",", ".")


def _set_table_row(table, row_index: int, values: list[object]) -> None:
    for cell, value in zip(table.rows[row_index].cells, values):
        cell.text = str(value)


def _build_report_assets(metadata: dict, output_dir: Path) -> dict[str, Path]:
    """Create report figures from the same data and metadata as the active app."""

    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np
    import seaborn as sns

    data = pd.read_csv(WORKSPACE_ROOT / "alzheimers_disease_data.csv")
    systolic = pd.to_numeric(data["SystolicBP"], errors="coerce")
    diastolic = pd.to_numeric(data["DiastolicBP"], errors="coerce")
    clean = data.loc[systolic > diastolic].copy()
    selected = list(metadata.get("selected_features", []))
    output_dir.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}

    def save(name: str, figure) -> None:
        path = output_dir / name
        figure.savefig(path, dpi=160, bbox_inches="tight")
        plt.close(figure)
        paths[name] = path

    counts = clean["Diagnosis"].astype(int).value_counts().reindex([0, 1], fill_value=0)
    figure, axis = plt.subplots(figsize=(7.5, 4.5))
    bars = axis.bar(["Không mắc (0)", "Mắc Alzheimer (1)"], counts.values, color=["#4C72B0", "#DD8452"])
    axis.set_title(f"Phân bố nhãn Diagnosis (n={len(clean):,}, sau lọc BP)".replace(",", "."))
    axis.set_ylabel("Số bệnh nhân")
    for bar, value in zip(bars, counts.values):
        axis.text(bar.get_x() + bar.get_width() / 2, bar.get_height(), f"{int(value):,}".replace(",", "."), ha="center", va="bottom")
    figure.tight_layout()
    save("report_diagnosis_distribution.png", figure)

    cohens = pd.DataFrame(metadata.get("eda", {}).get("cohens_d_top15", []))
    if not cohens.empty:
        cohens = cohens.sort_values("cohens_d")
        figure, axis = plt.subplots(figsize=(8.5, 6.0))
        values = cohens["cohens_d"].astype(float)
        axis.barh(cohens["feature"], values, color=["#4C72B0" if value < 0 else "#DD8452" for value in values])
        axis.axvline(0, color="#444444", linewidth=0.8)
        axis.set_title("Top 15 biến theo Cohen's d")
        axis.set_xlabel("Cohen's d")
        figure.tight_layout()
        save("report_cohens_d_top15.png", figure)

    correlation = pd.DataFrame(metadata.get("eda", {}).get("correlation", {}))
    if not correlation.empty:
        figure, axis = plt.subplots(figsize=(8.5, 7.0))
        sns.heatmap(correlation, cmap="coolwarm", center=0, annot=True, fmt=".2f", linewidths=0.2, ax=axis)
        axis.set_title("Ma trận tương quan các biến trong EDA artifact")
        figure.tight_layout()
        save("report_correlation_heatmap.png", figure)

    long_frame = clean[[*selected, "Diagnosis"]].melt(id_vars="Diagnosis", var_name="feature", value_name="value")
    figure, axes = plt.subplots(1, len(selected), figsize=(13.0, 3.8), constrained_layout=True)
    axes = np.atleast_1d(axes)
    for axis, feature in zip(axes, selected):
        sns.boxplot(data=long_frame[long_frame["feature"] == feature], x="Diagnosis", y="value", hue="Diagnosis", legend=False, ax=axis, palette=["#4C72B0", "#DD8452"])
        axis.set_title(feature)
        axis.set_xlabel("Diagnosis")
        axis.set_ylabel(feature)
    save("report_selected_feature_boxplots.png", figure)

    ranking = pd.DataFrame(metadata.get("feature_ranking", []))
    if not ranking.empty:
        top_ranking = ranking.head(15).sort_values("importance_mean")
        figure, axis = plt.subplots(figsize=(8.5, 6.0))
        axis.barh(top_ranking["feature"], top_ranking["importance_mean"].astype(float) * 100, color="#4C72B0")
        axis.set_title("Top 15 feature theo permutation importance trên validation")
        axis.set_xlabel("% đóng góp vào Accuracy")
        figure.tight_layout()
        save("report_feature_ranking_top15.png", figure)

    feature_selection = pd.DataFrame(metadata.get("feature_selection", []))
    if not feature_selection.empty:
        figure, left = plt.subplots(figsize=(8.5, 4.4))
        right = left.twinx()
        left.plot(feature_selection["feature_count"], feature_selection["accuracy_cv"], marker="o", color="#4C72B0", label="Accuracy")
        right.plot(feature_selection["feature_count"], feature_selection["pr_auc_cv"], marker="s", linestyle="--", color="#DD8452", label="PR-AUC")
        left.set_xlabel("Số lượng feature giữ lại")
        left.set_ylabel("Accuracy (validation)", color="#4C72B0")
        right.set_ylabel("PR-AUC (validation)", color="#DD8452")
        left.set_title("Accuracy / PR-AUC theo số lượng feature giữ lại")
        figure.tight_layout()
        save("report_feature_selection.png", figure)

    imbalance = pd.DataFrame(metadata.get("imbalance_comparison", []))
    if not imbalance.empty:
        figure, axes = plt.subplots(1, 2, figsize=(9.0, 3.8))
        methods = imbalance["method"].tolist()
        x = np.arange(len(methods))
        width = 0.18
        for offset, metric, color in zip([-1.5, -0.5, 0.5, 1.5], ["accuracy", "pr_auc", "f1", "recall"], ["#4C72B0", "#DD8452", "#55A868", "#C44E52"]):
            axes[0].bar(x + offset * width, imbalance[metric], width, label=metric.upper(), color=color)
        axes[0].set_xticks(x, methods, rotation=18)
        axes[0].set_ylim(0.80, 1.0)
        axes[0].set_title("Screening metrics")
        axes[0].legend(fontsize=7)
        axes[1].bar(x, imbalance["brier"], color="#8172B2")
        axes[1].set_xticks(x, methods, rotation=18)
        axes[1].set_title("Brier Score (thấp hơn = tốt hơn)")
        figure.tight_layout()
        save("report_imbalance_comparison.png", figure)

    benchmark = pd.DataFrame(metadata.get("benchmark", []))
    if not benchmark.empty:
        figure, axis = plt.subplots(figsize=(10.0, 4.8))
        x = np.arange(len(benchmark))
        width = 0.24
        axis.bar(x - width, benchmark["roc_auc_cv"], width, label="ROC-AUC", color="#4C72B0")
        axis.bar(x, benchmark["pr_auc_cv"], width, label="PR-AUC", color="#DD8452")
        axis.bar(x + width, benchmark["recall"], width, label="Recall", color="#55A868")
        labels = ["XGBoost" if name == "XGBoost" else name for name in benchmark["model"]]
        axis.set_xticks(x, labels, rotation=22, ha="right")
        axis.set_ylim(0, 1.0)
        axis.set_title("So sánh model bằng Repeated Stratified CV (5-fold x 2 lặp)")
        axis.legend()
        figure.tight_layout()
        save("report_benchmark.png", figure)

    calibration = metadata.get("calibration", {})
    before = calibration.get("before", {})
    after = calibration.get("after", {})
    if before and after and before.get("mean_predicted") and after.get("mean_predicted"):
        figure, axis = plt.subplots(figsize=(7.2, 4.8))
        axis.plot([0, 1], [0, 1], "k--", label="Calibraton hoàn hảo")
        axis.plot(before["mean_predicted"], before["fraction_positive"], "o-", label=f"Raw score (Brier={before['brier']:.4f})")
        axis.plot(after["mean_predicted"], after["fraction_positive"], "s-", label=f"Calibrated - Platt (Brier={after['brier']:.4f})")
        axis.set_xlabel("Mean predicted score")
        axis.set_ylabel("Observed positive rate")
        axis.set_title("Calibration Curve (validation set)")
        axis.legend(fontsize=8)
        figure.tight_layout()
        save("report_calibration.png", figure)

    confusion = np.asarray(metadata.get("test_confusion_matrix", {}).get("counts", []), dtype=float)
    if confusion.shape == (2, 2):
        figure, axis = plt.subplots(figsize=(4.8, 4.8))
        sns.heatmap(confusion, annot=True, fmt=".0f", cmap="Blues", cbar=False, ax=axis)
        axis.set_xlabel("Predicted label")
        axis.set_ylabel("True label")
        axis.set_xticklabels(["Không mắc", "Mắc bệnh"])
        axis.set_yticklabels(["Không mắc", "Mắc bệnh"], rotation=0)
        axis.set_title(f"Confusion Matrix - Test set (n={int(confusion.sum())})")
        figure.tight_layout()
        save("report_confusion_matrix.png", figure)

    try:
        import sys

        from sklearn.metrics import auc, roc_curve

        if str(ROOT) not in sys.path:
            sys.path.insert(0, str(ROOT))
        from ml_pipeline import _probabilities, load_artifacts

        payload, _ = load_artifacts(ROOT / "artifacts")
        locked = pd.read_csv(ROOT / "data" / "locked_test.csv")
        probabilities = _probabilities(payload["model"], locked[selected])
        false_positive_rate, true_positive_rate, _ = roc_curve(locked["Diagnosis"].astype(int), probabilities)
        figure, axis = plt.subplots(figsize=(6.8, 4.8))
        axis.plot(false_positive_rate, true_positive_rate, color="#DD8452", label=f"XGBoost (AUC={auc(false_positive_rate, true_positive_rate):.3f})")
        axis.plot([0, 1], [0, 1], "k--", label="Random (AUC=0.5)")
        axis.set_xlabel("False Positive Rate")
        axis.set_ylabel("True Positive Rate")
        axis.set_title("ROC Curve - Model cuối cùng (test set)")
        axis.legend()
        figure.tight_layout()
        save("report_roc_curve.png", figure)
    except Exception:
        pass

    if not ranking.empty:
        selected_ranking = ranking[ranking["feature"].isin(selected)].sort_values("importance_mean")
        figure, axis = plt.subplots(figsize=(7.5, 4.2))
        axis.barh(selected_ranking["feature"], selected_ranking["importance_mean"], color="#4C72B0")
        axis.set_title("Feature importance - Model cuối (5 feature)")
        axis.set_xlabel("Validation permutation importance")
        figure.tight_layout()
        save("report_feature_importance_selected.png", figure)

    return paths


def _replace_embedded_images(docx_path: Path, replacements: dict[str, Path]) -> None:
    """Replace image bytes in a saved DOCX while preserving its relationships."""

    temporary = docx_path.with_name(f".{docx_path.name}.tmp")
    try:
        with ZipFile(docx_path, "r") as source, ZipFile(temporary, "w") as target:
            for info in source.infolist():
                data = source.read(info.filename)
                replacement = replacements.get(info.filename)
                if replacement is not None:
                    data = replacement.read_bytes()
                target.writestr(info, data)
        os.replace(temporary, docx_path)
    finally:
        temporary.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_REPORT, help="DOCX template to refresh")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Output DOCX")
    args = parser.parse_args()
    # Keep caller-supplied relative paths relative.  On Windows this avoids
    # round-tripping a non-ASCII OneDrive folder through an ANSI code page.
    source_report = args.source
    output_report = args.output
    metadata = json.loads((ROOT / "artifacts" / "active_metadata.json").read_text(encoding="utf-8"))
    audit = pd.read_csv(ROOT / "artifacts" / "audit" / "cleaning_sensitivity.csv")
    profile = pd.read_csv(ROOT / "artifacts" / "audit" / "bp_removed_profile.csv")
    by_status = profile.set_index("status")
    by_strategy = audit.set_index("strategy")
    report = Document(source_report)

    def prevent_table_row_split() -> None:
        """Keep each table row intact when Word/LibreOffice paginates the report."""

        for table in report.tables:
            for row in table.rows:
                tr_pr = row._tr.get_or_add_trPr()
                if tr_pr.find(qn("w:cantSplit")) is None:
                    tr_pr.append(OxmlElement("w:cantSplit"))

    prevent_table_row_split()

    # Tables are the numerical source visible to the reader. Refresh every
    # model-dependent value so the DOCX cannot retain stale numbers from a
    # previous run of the pipeline.
    cleaning_table = report.tables[2]
    _set_table_row(cleaning_table, 1, ["Số dòng dữ liệu", _integer(metadata["rows_raw"])])
    _set_table_row(cleaning_table, 5, [
        "Sai quy tắc y khoa: SystolicBP ≤ DiastolicBP",
        f"{_integer(metadata['cleaning']['invalid_blood_pressure_rows_removed'])} / {_integer(metadata['rows_raw'])} ({metadata['cleaning']['invalid_blood_pressure_rows_removed'] / metadata['rows_raw'] * 100:.1f}%) — loại: BP không hợp lệ",
    ])
    _set_table_row(cleaning_table, 7, [
        "Số giá trị duy nhất của cột PatientID",
        f"{_integer(metadata['rows_raw'])} / {_integer(metadata['rows_raw'])} — là mã định danh, không phải đặc trưng dự đoán",
    ])

    cohens = metadata.get("eda", {}).get("cohens_d_top15", [])[:5]

    def effect_label(value: object) -> str:
        magnitude = abs(float(value))
        if magnitude >= 0.8:
            return "Lớn"
        if magnitude >= 0.5:
            return "Vừa"
        return "Nhỏ"

    for row_index, item in enumerate(cohens, start=1):
        signed = f"{float(item['cohens_d']):+.3f}".replace(".", ",")
        _set_table_row(report.tables[3], row_index, [item["feature"], signed, effect_label(item["cohens_d"])])

    feature_rows = metadata.get("feature_selection", [])
    _set_table_row(report.tables[4], 0, ["Số feature giữ lại", "Accuracy (Repeated 5-fold CV)", "PR-AUC (Repeated 5-fold CV)"])
    for row_index, item in enumerate(feature_rows, start=1):
        count = int(item["feature_count"])
        label = f"{count} (được chọn)" if count == len(metadata.get("selected_features", [])) else (f"{count} (đầy đủ)" if count == len(metadata.get("all_features", [])) else str(count))
        _set_table_row(report.tables[4], row_index, [label, _decimal(item["accuracy_cv"]), _decimal(item["pr_auc_cv"])])

    method_labels = {"Baseline": "Baseline (không xử lý)", "scale_pos_weight": "scale_pos_weight", "SMOTENC": "SMOTENC"}
    for row_index, item in enumerate(metadata.get("imbalance_comparison", []), start=1):
        _set_table_row(report.tables[5], row_index, [
            method_labels.get(item["method"], item["method"]),
            _decimal(item.get("accuracy")),
            _decimal(item.get("pr_auc")),
            _decimal(item.get("f1")),
            _decimal(item.get("recall")),
            _decimal(item.get("brier")),
        ])

    benchmark_by_model = {item["model"]: item for item in metadata.get("benchmark", [])}
    benchmark_labels = {
        "DummyClassifier": "DummyClassifier (base - lớp đa số)",
        "Logistic Regression": "Logistic Regression (base - học máy)",
        "KNN": "KNN (k=5)",
        "SVM (RBF)": "SVM (RBF)",
        "Random Forest": "Random Forest",
        "Gradient Boosting": "Gradient Boosting",
        "XGBoost": f"XGBoost ({metadata.get('selected_imbalance_method', 'scale_pos_weight')})",
    }
    benchmark_order = ["DummyClassifier", "Logistic Regression", "KNN", "SVM (RBF)", "Random Forest", "Gradient Boosting", "XGBoost"]
    for row_index, model_name in enumerate(benchmark_order, start=1):
        item = benchmark_by_model.get(model_name, {})
        _set_table_row(report.tables[7], row_index, [
            benchmark_labels[model_name],
            _decimal(item.get("pr_auc_cv")),
            _decimal(item.get("roc_auc_cv")),
            _decimal(item.get("f1")),
            _decimal(item.get("recall")),
        ])

    recall_order = ["KNN", "SVM (RBF)", "Logistic Regression", "Random Forest", "Gradient Boosting", "XGBoost"]
    for row_index, model_name in enumerate(recall_order, start=1):
        item = benchmark_by_model.get(model_name, {})
        _set_table_row(report.tables[8], row_index, [
            (
                f"XGBoost ({metadata.get('selected_imbalance_method', 'scale_pos_weight')})"
                if model_name == "XGBoost"
                else report.tables[8].rows[row_index].cells[0].text
            ),
            report.tables[8].rows[row_index].cells[1].text,
            _decimal(item.get("recall")),
            report.tables[8].rows[row_index].cells[3].text,
        ])

    params = metadata.get("hyperparameters", {})
    _set_table_row(report.tables[9], 1, ["learning_rate", _decimal(params.get("learning_rate"), 2)])
    _set_table_row(report.tables[9], 2, ["max_depth", int(params.get("max_depth", 0))])
    _set_table_row(report.tables[9], 3, ["n_estimators", int(params.get("n_estimators", 0))])
    tuning_pr_auc = metadata.get("tuning_best_pr_auc")
    _set_table_row(report.tables[9], 4, [
        "PR-AUC (5-fold CV) với tham số này",
        f"{_decimal(tuning_pr_auc)} (GridSearchCV, average_precision)" if tuning_pr_auc is not None else "Chưa lưu trong artifact",
    ])

    calibration = metadata.get("calibration", {})
    before = calibration.get("before", {})
    after = calibration.get("after", {})
    if before and after:
        _set_table_row(report.tables[10], 1, ["Brier Score", _decimal(before.get("brier")), _decimal(after.get("brier"))])
        _set_table_row(report.tables[10], 2, ["ECE (quantile 10 bins)", _decimal(before.get("ece")), _decimal(after.get("ece"))])

    test_metrics = metadata.get("test_metrics", {})
    ci = metadata.get("bootstrap_95_ci", {})
    metric_rows = [
        ("Accuracy", "accuracy", None),
        ("ROC-AUC", "roc_auc", "roc_auc"),
        ("PR-AUC", "pr_auc", "pr_auc"),
        ("F1", "f1", None),
        ("F2", "f2", None),
        ("Precision", "precision", "precision"),
        ("Recall", "recall", "recall"),
    ]
    for row_index, (label, metric_name, ci_name) in enumerate(metric_rows, start=1):
        ci_values = ci.get(ci_name) if ci_name else None
        ci_text = f"[{_decimal(ci_values[0], 4)} ; {_decimal(ci_values[1], 4)}]" if ci_values else "—"
        _set_table_row(report.tables[11], row_index, [label, _decimal(test_metrics.get(metric_name), 3), ci_text])
    _set_table_row(report.tables[12], 6, ["Vi phạm luật SystolicBP ≤ DiastolicBP", "Phát hiện, loại dòng"])

    toc_lines = (
        "15. Bổ sung phương pháp và tính minh bạch\t16",
        "15.1. Danh sách kiểm tra trước khi nộp\t16",
        "16. Tài liệu tham khảo\t17",
    )
    toc_present = all(any(paragraph.text == line for paragraph in report.paragraphs) for line in toc_lines)
    body_offset = 3 if toc_present else 0

    # A previous run inserted the new TOC lines before applying fixed body
    # indexes. Restore the original paragraphs that were consequently
    # overwritten, then use the offset for all subsequent refreshes.
    if toc_present:
        shifted_body_repairs = {
            43: "Phân tích khám phá dữ liệu (EDA)",
            45: "Chọn lọc đặc trưng (feature selection)",
            47: "So sánh xử lý mất cân bằng",
            51: "Tối ưu hyperparameter",
            69: "4. Bước 1 — Thu thập và làm sạch dữ liệu",
            90: "Hình 4. Boxplot của 5 biến phân biệt mạnh nhất, so sánh giữa nhóm không mắc (0) và mắc bệnh (1)",
            100: "",
            112: "8.2. Kết quả chạy:",
            115: "Kết luận model được giới hạn ở bảng CV và policy đã lưu; không dùng locked test để chọn model. XGBoost là họ model được giữ cho artifact V3, còn imbalance method là quyết định có kiểm soát và có thể thay đổi theo kết quả fold-level.",
            123: "",
            125: "",
            129: "",
            139: "Tab 1 — Sàng lọc 1 ca: nhập 5 chỉ số, trả về calibrated score và tín hiệu sàng lọc positive/negative; đây không phải chẩn đoán.",
            153: "3. Sao lưu (backup) dữ liệu hiện tại để có thể rollback nếu cần.",
            156: "Chấm champion và challenger trên cùng locked test; policy kiểm tra PR-AUC, recall, Brier, F2, paired bootstrap CI và hiện yêu cầu UI xác nhận trước khi promote. Active artifact/metadata chỉ được thay thế sau khi staging hai file thành công dưới file lock.",
            163: "15. Bổ sung phương pháp và tính minh bạch",
            164: "Phạm vi sử dụng và diễn giải lâm sàng. Mô hình trong đề tài nên được xem là công cụ sàng lọc và minh hoạ học máy trên dữ liệu bảng, không phải thiết bị chẩn đoán. Theo National Institute on Aging, đánh giá Alzheimer trong thực hành có thể bao gồm bệnh sử, đánh giá trí nhớ và tư duy, khám lâm sàng, xét nghiệm để loại trừ nguyên nhân khác, đánh giá tâm thần và trong một số trường hợp là hình ảnh hoặc biomarker. Vì vậy, kết quả của ứng dụng không được dùng làm căn cứ duy nhất cho quyết định y khoa. [1]",
            168: "Tính đồng bộ giữa báo cáo và mã nguồn. Bản aligned này được refresh từ active metadata và model artifact cùng release; app có validate_dataframe, calibration, locked-test evaluation và retrain loop. Bản nộp cần kèm đúng phiên bản mã nguồn, dataset và model artifacts tương ứng để tái lập các con số trong báo cáo. Việc báo cáo rõ mục tiêu sử dụng, dữ liệu, quy trình, hiệu năng, giới hạn và phiên bản là phù hợp với khuyến nghị TRIPOD+AI về tính minh bạch của nghiên cứu mô hình dự đoán. [5]",
            172: "• Đảm bảo báo cáo, app.py, notebook, dataset và model artifacts cùng mô tả một phiên bản pipeline.",
        }
        for index, text in shifted_body_repairs.items():
            paragraph = report.paragraphs[index + body_offset]
            paragraph.text = text
            if index in {43, 45, 47, 51}:
                for run in paragraph.runs:
                    run.bold = True
            elif index == 90:
                caption_reference = report.paragraphs[80].runs[0]
                target_rpr = paragraph.runs[0]._r.get_or_add_rPr()
                for child in list(target_rpr):
                    target_rpr.remove(child)
                for child in caption_reference._r.get_or_add_rPr():
                    target_rpr.append(deepcopy(child))
            elif index == 112:
                heading_reference = report.paragraphs[113].runs[0]
                target_rpr = paragraph.runs[0]._r.get_or_add_rPr()
                for child in list(target_rpr):
                    target_rpr.remove(child)
                for child in heading_reference._r.get_or_add_rPr():
                    target_rpr.append(deepcopy(child))

        # The first repair run cleared the drawing paragraph immediately
        # before the calibration caption. Restore the original embedded chart
        # from the DOCX media package and keep its original visual scale.
        calibration_paragraph = report.paragraphs[123 + body_offset]
        if "w:drawing" not in calibration_paragraph._p.xml:
            with ZipFile(source_report) as source_docx:
                calibration_png = source_docx.read("word/media/image12.png")
            calibration_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            calibration_paragraph.add_run().add_picture(BytesIO(calibration_png), width=Inches(4.0))

    def set_paragraph(index: int, text: str) -> None:
        report.paragraphs[index + body_offset].text = text

    set_paragraph(
        46,
        "Permutation importance được tính trên validation để audit. Chế độ mặc định report_v3 giữ cố định 5/32 feature chính thức nhằm tái lập báo cáo; chế độ nested_cv được cung cấp riêng để đo độ ổn định lựa chọn trong inner folds, không xếp hạng trên outer-test.",
    )
    set_paragraph(
        48,
        "Baseline, scale_pos_weight và SMOTENC được chấm trên cùng 10 folds. Policy chọn method theo PR-AUC trung bình; trong vùng chênh dưới 0,005 ưu tiên phương pháp đơn giản hơn, đồng thời lưu SD và số fold.",
    )
    set_paragraph(
        50,
        "Base model + 5 ứng viên, Repeated Stratified CV (5-fold x 2 lặp); bảng imbalance lưu metric từng fold/tóm tắt mean ± SD và không dùng locked test để chọn.",
    )
    set_paragraph(
        54,
        f"Platt scaling + ngưỡng tối ưu F2 trên validation; threshold hiện tại của active artifact là {metadata['threshold']:.3f}. Score chỉ là screening score đã hiệu chỉnh trong cohort nội bộ, không phải xác suất lâm sàng đã xác nhận bên ngoài.",
    )
    set_paragraph(
        86,
        "Ma trận tương quan được lấy từ EDA artifact cho các biến đại diện và Diagnosis; kết quả dùng để mô tả quan hệ tuyến tính, không phải bằng chứng nhân quả hay bảo đảm không có multicollinearity ở mọi cohort.",
    )
    removed = by_status.loc["removed"]
    kept = by_status.loc["kept"]
    remove = by_strategy.loc["remove"]
    swap = by_strategy.loc["swap-if-plausible"]
    keep = by_strategy.loc["keep-with-flag"]
    set_paragraph(
        72,
        f"Kết luận & xử lý: dữ liệu thô có 2.149 dòng và 35 cột; có {int(removed['rows'])} dòng BP không hợp lệ ({removed['row_pct']:.2f}%). Tỷ lệ Diagnosis của nhóm removed là {removed['diagnosis_rate']:.3f}, nhóm kept là {kept['diagnosis_rate']:.3f}; profile trung vị, IQR và SMD được lưu trong artifacts/audit/bp_removed_profile.csv. Theo sensitivity trên cùng 393 dòng đánh giá nội bộ: remove cho PR-AUC {remove['pr_auc_locked']:.3f}, recall {remove['recall_locked']:.3f}, F2 {remove['f2_locked']:.3f}, Brier {remove['brier_locked']:.3f}; swap-if-plausible cho PR-AUC {swap['pr_auc_locked']:.3f}, recall {swap['recall_locked']:.3f}, F2 {swap['f2_locked']:.3f}, Brier {swap['brier_locked']:.3f}; keep-with-flag cho PR-AUC {keep['pr_auc_locked']:.3f}, recall {keep['recall_locked']:.3f}, F2 {keep['f2_locked']:.3f}, Brier {keep['brier_locked']:.3f}. Vì ba chiến lược cho trade-off khác nhau, pipeline chính giữ quy tắc remove để tái lập V3 và lưu audit, không âm thầm sửa BP.",
    )
    set_paragraph(
        93,
        "Chế độ report_v3 giữ 5 feature (MMSE, FunctionalAssessment, ADL, MemoryComplaints, BehavioralProblems) cho khả năng tái lập. Đây là feature set chính thức cố định, không nên mô tả là lựa chọn hoàn toàn tự động; nested_cv có thể chạy riêng để lưu feature stability và outer-fold uncertainty.",
    )
    set_paragraph(
        103,
        f"Baseline, SMOTENC và scale_pos_weight được chấm trên cùng splits. Active artifact chọn {metadata['selected_imbalance_method']} theo policy: PR-AUC mean cao nhất là {metadata['imbalance_selection']['best_mean_pr_auc']:.4f}, nhưng các method trong khoảng 0,005 được ưu tiên theo độ đơn giản. Vì vậy không tuyên bố một method thắng mọi chỉ số; bảng metadata lưu mean, SD và fold count.",
    )
    set_paragraph(
        115,
        "Kết luận model được giới hạn ở bảng CV và policy đã lưu; không dùng locked test để chọn model. XGBoost là họ model được giữ cho artifact V3, còn imbalance method là quyết định có kiểm soát và có thể thay đổi theo kết quả fold-level.",
    )
    set_paragraph(
        118,
        f"Kết luận về mất cân bằng được diễn giải theo mean ± SD trên cùng folds, không theo một con số đơn lẻ. Active metadata hiện chọn {metadata['selected_imbalance_method']} theo policy; PR-AUC tốt nhất trong vùng so sánh là {metadata['imbalance_selection']['best_mean_pr_auc']:.4f}, còn method được chọn có F2 {metadata['imbalance_selection'].get('selected_mean_f2', float('nan')):.4f} và recall {metadata['imbalance_selection'].get('selected_mean_recall', float('nan')):.4f}. Threshold/calibration vẫn được chọn riêng trên validation.",
    )
    set_paragraph(
        126,
        f"Calibration summary trên validation ghi Brier {metadata['calibration']['validation_summary']['brier']:.4f}, ECE {metadata['calibration']['validation_summary'].get('ece', float('nan')):.4f}, prevalence {metadata['calibration']['validation_summary']['prevalence']:.3f} và null Brier {metadata['calibration']['validation_summary']['null_brier']:.4f}. Trước calibration, Brier/ECE là {metadata['calibration'].get('before', {}).get('brier', float('nan')):.4f}/{metadata['calibration'].get('before', {}).get('ece', float('nan')):.4f}; sau Platt là {metadata['calibration'].get('after', {}).get('brier', float('nan')):.4f}/{metadata['calibration'].get('after', {}).get('ece', float('nan')):.4f}. Bảng threshold 0,05–0,95 gồm precision, recall, specificity, F1, F2 và confusion counts được lưu trong metadata; threshold {metadata['threshold']:.3f} được chọn bằng F2 trên validation với tie-break gần 0,490.",
    )
    set_paragraph(
        128,
        f"Fit trên {metadata['split_sizes']['train_validation']:,} dòng train+validation sau khi loại locked cohort, đánh giá đúng 1 lần trên locked internal Test ({metadata['split_sizes']['locked_test']:,} dòng; hash {metadata['locked_test_sha256'][:12]}…). Locked test không được dùng để chọn model, threshold hoặc calibration.",
    )
    set_paragraph(
        132,
        f"Confusion matrix: TN={int(metadata['test_metrics']['true_negatives'])}, FP={int(metadata['test_metrics']['false_positives'])}, FN={int(metadata['test_metrics']['false_negatives'])}, TP={int(metadata['test_metrics']['true_positives'])}. Trên locked internal test: PR-AUC={metadata['test_metrics']['pr_auc']:.3f}, ROC-AUC={metadata['test_metrics']['roc_auc']:.3f}, recall={metadata['test_metrics']['recall']:.3f}, F2={metadata['test_metrics']['f2']:.3f}, Brier={metadata['test_metrics']['brier']:.3f}. Đây là tín hiệu sàng lọc nội bộ, không phải chẩn đoán.",
    )
    set_paragraph(
        142,
        f"Tab Model information hiển thị metric, bootstrap CI, hyperparameter, threshold, feature list, cohort, subgroup metrics có n/CI và external validation status. Active release hiện là {metadata['version']} với model SHA-256 {metadata['model_sha256'][:16]}…; báo cáo này được refresh từ cùng active metadata/model artifact.",
    )
    set_paragraph(
        156,
        "Chấm champion và challenger trên cùng locked test; policy kiểm tra PR-AUC, recall, Brier, F2, paired bootstrap CI và yêu cầu UI xác nhận trước khi promote. Khi promote thành công, code ghi lại dataset đã gộp sau khi tạo backup; nếu policy không đạt hoặc chưa xác nhận, dataset và active artifact được giữ nguyên.",
    )
    set_paragraph(
        136,
        "Hình 12. Permutation importance của 5 feature trong model cuối trên validation",
    )
    set_paragraph(
        121,
        f"GridSearchCV lặp qua 27 tổ hợp tham số x 5-fold = 135 lần huấn luyện, chọn tổ hợp có PR-AUC trung bình cao nhất. Với artifact hiện tại, PR-AUC tốt nhất của GridSearchCV là {metadata.get('tuning_best_pr_auc', float('nan')):.4f} trước khi refit model cuối.",
    )
    set_paragraph(
        139,
        "Ứng dụng Streamlit gồm 5 mục đúng với implementation: Batch CSV + Data processing (validate_dataframe, imputation, rejected-row export và monitoring), EDA, Model information, Single case (5 input features) và Retrain / Model update (preview challenger, policy và xác nhận promote). Tất cả đều dùng cùng active artifact và schema 5 feature; kết quả chỉ là screening, không phải chẩn đoán.",
    )
    set_paragraph(
        159,
        "Đề tài đã hoàn thành một pipeline Machine Learning end-to-end có kiểm soát: làm sạch và audit dữ liệu; EDA; feature set report_v3 cố định; so sánh imbalance có SD/fold count; tuning, calibration và threshold từ validation; locked internal evaluation; versioned artifact; subgroup analysis; và ứng dụng Streamlit phục vụ học tập.",
    )
    set_paragraph(
        166,
        f"Hiệu chỉnh xác suất và ngưỡng: threshold {metadata['threshold']:.3f} được chọn trên validation bằng F2; bảng operating points được lưu để kiểm tra precision/recall/specificity/FN. Brier và calibration curve chỉ phản ánh cohort nội bộ; score không nên mặc định được hiểu là xác suất bệnh ở quần thể hoặc cơ sở y tế khác.",
    )
    set_paragraph(
        167,
        "Subgroup metrics theo Gender, nhóm tuổi, Ethnicity, EducationLevel, Diabetes, Hypertension và FamilyHistoryAlzheimers được lưu kèm n, CI và trạng thái estimable/not estimable trong metadata. External validation hiện là not performed; cần dữ liệu ngoài trước khi diễn giải theo hướng lâm sàng.",
    )
    set_paragraph(
        171,
        "Cập nhật mục lục/trang sau khi render lại Word; artifact active, report V3 và audit tables phải cùng release.",
    )
    set_paragraph(
        175,
        "Đã bổ sung subgroup metrics có cỡ mẫu/CI; external validation vẫn ghi rõ not performed và có script evaluate_external.py để chấm dữ liệu ngoài mà không fit lại model.",
    )

    # The report's table of contents is intentionally static so the submitted
    # DOCX remains readable in viewers that do not recalculate fields.
    report.paragraphs[15].text = "5.4. So sánh phân bố theo nhóm chẩn đoán\t7"
    if not toc_present:
        toc_anchor = report.paragraphs[25]
        toc_style = report.paragraphs[24].style
        for toc_line in toc_lines:
            inserted = toc_anchor.insert_paragraph_before(toc_line)
            inserted.style = toc_style
    output_report.parent.mkdir(parents=True, exist_ok=True)
    report.save(output_report)
    report_assets = _build_report_assets(metadata, DELIVERABLES_ROOT / "assets" / "report_assets_aligned")
    media_map = {
        "word/media/image4.png": "report_diagnosis_distribution.png",
        "word/media/image5.png": "report_cohens_d_top15.png",
        "word/media/image6.png": "report_correlation_heatmap.png",
        "word/media/image7.png": "report_selected_feature_boxplots.png",
        "word/media/image8.png": "report_feature_ranking_top15.png",
        "word/media/image9.png": "report_feature_selection.png",
        "word/media/image10.png": "report_imbalance_comparison.png",
        "word/media/image11.png": "report_benchmark.png",
        "word/media/image12.png": "report_calibration.png",
        "word/media/image13.png": "report_confusion_matrix.png",
        "word/media/image14.png": "report_roc_curve.png",
        "word/media/image15.png": "report_feature_importance_selected.png",
    }
    _replace_embedded_images(
        output_report,
        {media_name: report_assets[asset_name] for media_name, asset_name in media_map.items()},
    )
    print(output_report)


if __name__ == "__main__":
    main()
